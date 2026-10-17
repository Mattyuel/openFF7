"""Honey Bee scene checklist: python3 tools/mk_scene_checklist.py docs/scene_checklist.xlsx (needs openpyxl)."""
import sys, re, importlib
sys.path[:0] = ["patches", "tools"]
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from _honey_bee_data import GIRLS
import _honey_bee_scenes_more as more
hb = importlib.import_module("620_honey_bee")

POS = {"mission": "Mission (she lies on her back, Cloud on top)", "cowgirl": "Cowgirl (Cloud on his back, she kneels astride facing him)",
       "ontop": "Ontop (as cowgirl, she faces the other way)", "dog": "Dog (she's on all fours, Cloud kneels behind)",
       "head": "Head (both on the floor left of the bed: Cloud standing facing left, she kneels facing him)",
       "edge": "Edge (she stands bent over the bed's left end, Cloud standing behind her)",
       "headrev": "Head reverse (she lies on her back, head hanging off the bed's left end, Cloud standing at it)",
       "beside": "Beside (both lying on the bed)", "sit": "Sit (both sitting on the bed's edge)", "lie": "Lie"}
MUSIC = {5: "Fanfare (victory)", 8: "Fiddle de Chocobo", 16: "Rufus' Welcoming Ceremony", 38: "Crazy Motorcycle",
         46: "Chocobo fanfare", 49: "Interrupted by Fireworks", 62: "Debut", 72: "Steal the Tiny Bronco! (unconfirmed track)"}
ROOM = {"Basic": "convil_3 (normal room)", "Close": "onna_3 (close room)", "Devoted": "onna_6 (devoted room, dimmed)"}

def clean(t):
    t = t.replace("{NEW}", " / ").replace("{CLOUD}\n", "").replace("\n", " ")
    return re.sub(r"\s+", " ", t).strip()

def steps(sc):
    out = []
    for st in sc:
        k, v = st[0], (st[1] if len(st) > 1 else "")
        if k == "shake": out.append("[screen shake]"); continue
        if k == "G": out.append("Her: " + clean(v))
        elif k == "C": out.append("Cloud: " + clean(v))
        elif k == "g": out.append("[she: %s]" % v)
        elif k == "c": out.append("[Cloud: %s]" % v)
        elif k == "bed": out.append("[FADE -> %s]" % v)
        elif k == "start": out.append("[starts: %s]" % v)
        else: out.append("[%s %s]" % (k, v))
    return "\n".join(out)

rows = []
for tier, scenes, aff in (("Basic", None, "0-9"), ("Close", more.CLOSE, "10-19"), ("Devoted", more.DEVOTED, "20")):
    for i, g in enumerate(GIRLS):
        sc = g["scene"] if scenes is None else scenes[g["name"]]
        lay = [st[1] for st in sc if st[0] == "bed"]
        pos = POS.get(lay[0], lay[0]) if lay else "(no bed part)"
        if tier == "Devoted":
            m = hb.PART2_MUSIC_GIRL.get(g["name"], hb.PART2_MUSIC.get(lay[0] if lay else "", hb.PART2_MUSIC_DEFAULT))
            music = "Part 2: %s; closing text: Good Night, Until Tomorrow!" % MUSIC.get(m, m)
        else:
            music = "Closing text: Good Night, Until Tomorrow!"
        rows.append((g["name"], "*" * g["stars"], tier, aff, ROOM[tier], pos, music, steps(sc)))
order = {g["name"]: k for k, g in enumerate(GIRLS)}
rows.sort(key=lambda r: (order[r[0]], ("Basic", "Close", "Devoted").index(r[2])))

wb = Workbook(); ws = wb.active; ws.title = "Scenes"
hdr = ["#", "Girl", "Stars", "Tier", "Affection", "Room", "Position (after the fade)", "Music", "Scene (in order)",
       "Loads + finishes", "Positions", "Animations", "Text", "Music OK", "Notes"]
ws.append(hdr)
for n, r in enumerate(rows, 1):
    ws.append([n, *r, "Not tested", "Not tested", "Not tested", "Not tested", "Not tested", ""])
widths = [4, 9, 6, 8, 9, 20, 30, 30, 70, 12, 12, 12, 12, 12, 45]
for c, w in enumerate(widths, 1):
    ws.column_dimensions[ws.cell(1, c).column_letter].width = w
thin = Side(style="thin", color="BFBFBF")
tier_fill = {"Basic": "F2F2F2", "Close": "DDEBF7", "Devoted": "FCE4D6"}
N = len(rows) + 1
for row in ws.iter_rows(min_row=1, max_row=N):
    for cell in row:
        cell.font = Font(name="Arial", size=10, bold=cell.row == 1, color="FFFFFF" if cell.row == 1 else "000000")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
        if cell.row == 1:
            cell.fill = PatternFill("solid", fgColor="305496")
    if row[0].row > 1:
        for c in range(0, 9):
            row[c].fill = PatternFill("solid", fgColor=tier_fill[row[3].value])
        for c in range(9, 15):
            row[c].fill = PatternFill("solid", fgColor="FFF2CC")
ws.freeze_panes = "C2"
ws.auto_filter.ref = "A1:O%d" % N
dv = DataValidation(type="list", formula1='"Not tested,Pass,Fail,Partly"', allow_blank=False)
ws.add_data_validation(dv); dv.add("J2:N%d" % N)
for val, col in (("Pass", "C6EFCE"), ("Fail", "FFC7CE"), ("Partly", "FFEB9C")):
    ws.conditional_formatting.add("J2:N%d" % N, CellIsRule(operator="equal", formula=['"%s"' % val],
                                                            fill=PatternFill("solid", fgColor=col)))

s = wb.create_sheet("How to")
lines = [
    ("openFF7: Honey Bee scene checklist (60 scenes)", True),
    ("Each girl has three scenes, chosen by her affection when you pick her in the manager's Visit menu:", False),
    ("  Basic = affection 0-9 (convil_3), Close = 10-19 (onna_3), Devoted = 20 (onna_6, dimmed, bare-legged Cloud in part 2).", False),
    ("Debug build (L1+R1): Honey Bee -> Job + hire all; set All affection 0 / 10 / 20; Travel -> Honey Bee Inn; talk to the manager.", False),
    ("Visit menu: One Star Girls .. Five Star Girls pages (Next / Previous page); unhired girls show as ????.", False),
    ("Fill in the yellow columns: Loads + finishes (no black screen / freeze, back in the lobby), Positions (where people stand, lie, kneel),", False),
    ("Animations (poses, speed, nothing stuck), Text (fits the window, reads right), Music OK. Details in Notes.", False),
    ("Scene column: [FADE -> position] is the cut to part 2; [she: x] / [Cloud: x] are poses or moves.", False),
    ("", False),
    ("Totals", True),
]
for t, b in lines:
    s.append([t]); s.cell(s.max_row, 1).font = Font(name="Arial", size=14 if b and s.max_row == 1 else 10, bold=b)
for k, col in enumerate("JKLMN"):
    s.append([hdr[9 + k]] + ['=COUNTIF(Scenes!%s2:%s%d,"%s")' % (col, col, N, v) for v in ("Pass", "Fail", "Partly", "Not tested")])
s.append(["", "Pass", "Fail", "Partly", "Not tested"])
s.move_range("A%d:E%d" % (s.max_row, s.max_row), rows=-6)
s.column_dimensions["A"].width = 22
for c in "BCDE":
    s.column_dimensions[c].width = 12
wb.save(sys.argv[1])
print(len(rows), "rows")
